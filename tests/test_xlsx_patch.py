"""xlsx edits patch the package instead of rebuilding it: nothing unedited is lost."""

from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Font
from PIL import Image as PilImage

from app.services.file import xlsx_package, xlsx_patch


def workbook_with_chart_and_image() -> bytes:
    book = openpyxl.Workbook()
    data = book.active
    data.title = "Data"
    for row in ([1, 10], [2, 20], [3, 30]):
        data.append(row)
    data["C1"] = "=A1+B1"
    data["A1"].font = Font(bold=True)
    data.add_chart(BarChart(), "E2")
    data._charts[0].add_data(Reference(data, min_col=2, min_row=1, max_row=3))
    picture = io.BytesIO()
    PilImage.new("RGB", (4, 4), "red").save(picture, "PNG")
    data.add_image(Image(picture), "H2")
    book.create_sheet("Notes")["A1"] = "keep me"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def names(content: bytes) -> set[str]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return set(archive.namelist())


def media_and_charts(content: bytes) -> set[str]:
    return {n for n in names(content) if n.startswith(("xl/charts/", "xl/media/", "xl/drawings/"))}


def test_cell_edit_keeps_charts_images_formulas_styles_and_other_sheets():
    original = workbook_with_chart_and_image()
    assert any("chart" in n for n in media_and_charts(original))

    edited = xlsx_package.edit_xlsx(original, edits={"Data": {(2, 2): 99, (4, 1): "new row"}})

    assert media_and_charts(edited) == media_and_charts(original)
    book = openpyxl.load_workbook(io.BytesIO(edited))
    data = book["Data"]
    assert data["B2"].value == 99 and data["A4"].value == "new row"
    assert data["C1"].value == "=A1+B1"  # untouched formula
    assert data["A1"].font.bold is True  # untouched style
    assert data["B1"].value == 10 and data["B3"].value == 30
    assert book["Notes"]["A1"].value == "keep me"
    assert len(data._charts) == 1 and len(data._images) == 1


def test_overwriting_a_cell_keeps_its_style_and_clearing_keeps_the_cell():
    edited = xlsx_package.edit_xlsx(
        workbook_with_chart_and_image(), edits={"Data": {(1, 1): "text", (3, 2): None}}
    )
    data = openpyxl.load_workbook(io.BytesIO(edited))["Data"]
    assert data["A1"].value == "text" and data["A1"].font.bold is True
    assert data["B3"].value is None


def test_values_are_xml_escaped_and_typed():
    edited = xlsx_package.edit_xlsx(
        workbook_with_chart_and_image(),
        edits={"Notes": {(1, 1): "a<b>&\x01c", (1, 2): True, (1, 3): 1.5, (1, 4): 7}},
    )
    notes = openpyxl.load_workbook(io.BytesIO(edited))["Notes"]
    assert [c.value for c in notes[1]] == ["a<b>&c", True, 1.5, 7]


def test_new_sheet_is_added_and_the_package_stays_loadable():
    original = workbook_with_chart_and_image()
    edited = xlsx_package.edit_xlsx(original, new_sheets={"Fresh": [["x", 1], [None, 2]]})

    book = openpyxl.load_workbook(io.BytesIO(edited))
    assert book.sheetnames == ["Data", "Notes", "Fresh"]
    assert [[c.value for c in row] for row in book["Fresh"]] == [["x", 1], [None, 2]]
    assert media_and_charts(edited) == media_and_charts(original)


def test_deleting_a_sheet_removes_only_that_sheet():
    original = workbook_with_chart_and_image()
    edited = xlsx_package.edit_xlsx(original, delete="Notes")

    book = openpyxl.load_workbook(io.BytesIO(edited))
    assert book.sheetnames == ["Data"]
    assert len(book["Data"]._charts) == 1


def test_calc_chain_is_dropped_and_recalculation_requested():
    book = openpyxl.load_workbook(io.BytesIO(workbook_with_chart_and_image()))
    output = io.BytesIO()
    book.save(output)
    edited = xlsx_package.edit_xlsx(output.getvalue(), edits={"Data": {(1, 2): 5}})
    with zipfile.ZipFile(io.BytesIO(edited)) as archive:
        assert "xl/calcChain.xml" not in archive.namelist()
        assert 'fullCalcOnLoad="1"' in archive.read("xl/workbook.xml").decode()


def test_overwriting_a_shared_formula_anchor_is_refused():
    sheet = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        '<row r="1"><c r="A1"><f t="shared" ref="A1:A2" si="0">B1</f><v>1</v></c></row>'
        "</sheetData></worksheet>"
    )
    with pytest.raises(ValueError, match="shared or array formula"):
        xlsx_patch.patch_sheet_xml(sheet, {(1, 1): 5})


def test_rejects_non_workbooks_and_unknown_sheets():
    with pytest.raises(ValueError, match=r"Not a valid .xlsx"):
        xlsx_package.edit_xlsx(b"not a zip", edits={})
    with pytest.raises(ValueError, match="Sheet 'Nope' not found"):
        xlsx_package.edit_xlsx(workbook_with_chart_and_image(), edits={"Nope": {(1, 1): 1}})
    with pytest.raises(ValueError, match="expands to more than"):
        xlsx_package.edit_xlsx(workbook_with_chart_and_image(), edits={}, max_unpacked_bytes=10)


def test_documents_with_a_dtd_are_refused():
    buffer = io.BytesIO(workbook_with_chart_and_image())
    hostile = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(hostile, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/workbook.xml":
                data = b'<!DOCTYPE x [<!ENTITY a "aaaa">]>' + data
            target.writestr(name, data)
    with pytest.raises(ValueError, match="DTD"):
        xlsx_package.edit_xlsx(hostile.getvalue(), edits={})
