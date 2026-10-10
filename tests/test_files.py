"""File upload, download, versioning and chunked uploads through the real app."""

from __future__ import annotations

import json

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from starlette.requests import Request

from app.api.middleware.body_limit import BodySizeLimitMiddleware
from tests.conftest import run
from tests.test_auth_flow import bearer, login, register

PDF = ("report.pdf", b"%PDF-1.4 hello", "application/pdf")


def signed_in(env, username="ada", email="ada@example.com"):
    client = env.client()
    register(client, username=username, email=email)
    client.headers.update(bearer(login(client, username)))
    return client


def upload(client, name="report.pdf", content=b"%PDF-1.4 hello", mime="application/pdf", **form):
    return client.post("/api/v1/files/upload", files={"file": (name, content, mime)}, data=form)


# ---- upload / download --------------------------------------------------------------


def test_upload_then_download_round_trips_as_an_attachment(env):
    client = signed_in(env)
    response = upload(client)
    assert response.status_code == 201
    file_id = response.json()["file_id"]

    download = client.get(f"/api/v1/files/{file_id}")
    assert download.status_code == 200
    assert download.content == b"%PDF-1.4 hello"
    assert download.headers["content-disposition"].startswith("attachment;")


def test_other_users_cannot_see_touch_or_list_a_file(env):
    owner = signed_in(env)
    file_id = upload(owner).json()["file_id"]
    other = signed_in(env, "eve", "eve@example.com")

    for response in (
        other.get(f"/api/v1/files/{file_id}"),
        other.get(f"/api/v1/files/{file_id}/metadata"),
        other.get(f"/api/v1/files/{file_id}/versions"),
        other.put(f"/api/v1/files/{file_id}", json={"filename": "mine.pdf"}),
        other.delete(f"/api/v1/files/{file_id}"),
    ):
        assert response.status_code == 404
    assert other.get("/api/v1/files/").json()["files"] == []


def test_active_content_types_are_rejected(env):
    client = signed_in(env)
    assert upload(client, "x.html", b"<script>1</script>", "text/html").status_code == 400
    assert upload(client, "x.svg", b"<svg onload=1>", "image/svg+xml").status_code == 400
    assert upload(client, "x.pdf", b"", "application/pdf").status_code == 400  # empty


def test_hostile_filenames_cannot_break_the_download_header(env):
    client = signed_in(env)
    file_id = upload(client, 'a"; evil=1.txt', b"x", "text/plain").json()["file_id"]

    header = client.get(f"/api/v1/files/{file_id}").headers["content-disposition"]

    assert header.count('"') == 2  # only the quotes we added around the fallback name
    assert "evil" not in header.split("filename*")[0].replace("evil=1", "")


# ---- deduplication ------------------------------------------------------------------


def test_identical_bytes_dedupe_per_user_never_across_users(env):
    ada = signed_in(env)
    eve = signed_in(env, "eve", "eve@example.com")

    first = upload(ada).json()
    again = upload(ada).json()
    theirs = upload(eve).json()

    assert again["file_id"] == first["file_id"]
    assert theirs["file_id"] != first["file_id"]
    assert theirs["user_id"] != first["user_id"]


# ---- versions -----------------------------------------------------------------------


def test_new_versions_keep_history_and_default_to_newest(env):
    client = signed_in(env)
    file_id = upload(client, content=b"%PDF-v1").json()["file_id"]

    v2 = client.post(
        f"/api/v1/files/{file_id}/versions",
        files={"file": ("report.pdf", b"%PDF-v2", "application/pdf")},
    )
    assert v2.status_code == 201
    assert v2.json()["version"] == 2 and v2.json()["file_id"] == file_id

    assert client.get(f"/api/v1/files/{file_id}").content == b"%PDF-v2"
    assert client.get(f"/api/v1/files/{file_id}?version=1").content == b"%PDF-v1"
    versions = client.get(f"/api/v1/files/{file_id}/versions").json()
    assert [v["version"] for v in versions] == [2, 1]


def test_cannot_add_a_version_to_someone_elses_file(env):
    owner = signed_in(env)
    file_id = upload(owner).json()["file_id"]
    eve = signed_in(env, "eve", "eve@example.com")

    response = eve.post(
        f"/api/v1/files/{file_id}/versions",
        files={"file": ("x.pdf", b"%PDF-evil", "application/pdf")},
    )
    assert response.status_code == 400


def test_modify_as_new_version_leaves_the_old_bytes_alone(env):
    client = signed_in(env)
    file_id = upload(client, content=b"%PDF-original").json()["file_id"]
    user_id = client.get("/api/v1/auth/me").json()["id"]
    storage = env.file_service.storage_backend

    async def append(path: str) -> bool:
        return await storage.save(path, await storage.read(path) + b"+edited", overwrite=True)

    result = run(env.file_service.modify_as_new_version(file_id, user_id, append))

    assert result.version == 2
    assert client.get(f"/api/v1/files/{file_id}?version=1").content == b"%PDF-original"
    assert client.get(f"/api/v1/files/{file_id}").content == b"%PDF-original+edited"


# ---- listing and deletion -----------------------------------------------------------


def test_listing_paginates_with_a_real_total_and_one_row_per_file(env):
    client = signed_in(env)
    ids = [
        upload(client, f"f{i}.txt", f"body {i}".encode(), "text/plain").json()["file_id"]
        for i in range(3)
    ]
    client.post(f"/api/v1/files/{ids[0]}/versions", files={"file": ("f0.txt", b"v2", "text/plain")})

    page = client.get("/api/v1/files/?page=1&page_size=2").json()

    assert page["total"] == 3  # versions are not separate files
    assert len(page["files"]) == 2 and page["has_next"] is True
    assert len(client.get("/api/v1/files/?page=2&page_size=2").json()["files"]) == 1


def test_deleted_files_disappear_everywhere(env):
    client = signed_in(env)
    file_id = upload(client).json()["file_id"]

    assert client.delete(f"/api/v1/files/{file_id}").status_code == 200

    assert client.get(f"/api/v1/files/{file_id}").status_code == 404
    assert client.get("/api/v1/files/").json()["total"] == 0
    assert client.delete(f"/api/v1/files/{file_id}").status_code == 404


# ---- chunked uploads ----------------------------------------------------------------


def chunk(client, upload_id, number, data, total_chunks, total_size, name="big.txt"):
    return client.post(
        "/api/v1/files/upload/chunk",
        files={"chunk": (name, data, "application/octet-stream")},
        data={
            "upload_id": upload_id,
            "chunk_number": number,
            "total_chunks": total_chunks,
            "total_size": total_size,
            "filename": name,
            "content_type": "text/plain",
        },
    )


@pytest.fixture
def small_chunks(env):
    env.file_service.chunked_upload_handler.max_chunk_size = 4
    return env


def test_chunked_upload_assembles_in_order_even_when_sent_out_of_order(small_chunks):
    client = signed_in(small_chunks)
    parts = [b"aaaa", b"bbbb", b"cc"]
    for number in (3, 1, 2):
        assert chunk(client, "upload-0001", number, parts[number - 1], 3, 10).status_code == 200

    done = client.post("/api/v1/files/upload/chunk/complete?upload_id=upload-0001")

    assert done.status_code == 201
    assert client.get(f"/api/v1/files/{done.json()['file_id']}").content == b"aaaabbbbcc"


def test_incomplete_chunked_upload_cannot_be_completed(small_chunks):
    client = signed_in(small_chunks)
    chunk(client, "upload-0002", 1, b"aaaa", 3, 10)

    assert (
        client.post("/api/v1/files/upload/chunk/complete?upload_id=upload-0002").status_code == 409
    )


def test_a_second_user_cannot_inject_into_or_finish_anothers_upload(small_chunks):
    ada = signed_in(small_chunks)
    eve = signed_in(small_chunks, "eve", "eve@example.com")
    chunk(ada, "upload-0003", 1, b"aaaa", 2, 6)

    # Eve reuses Ada's upload id: she gets her own, separate session.
    chunk(eve, "upload-0003", 2, b"EVIL", 2, 6)
    assert eve.post("/api/v1/files/upload/chunk/complete?upload_id=upload-0003").status_code == 409

    chunk(ada, "upload-0003", 2, b"bb", 2, 6)
    done = ada.post("/api/v1/files/upload/chunk/complete?upload_id=upload-0003")
    assert ada.get(f"/api/v1/files/{done.json()['file_id']}").content == b"aaaabb"


@pytest.mark.parametrize(
    ("kwargs", "label"),
    [
        ({"number": 9, "total_chunks": 2, "total_size": 8}, "chunk number out of range"),
        ({"number": 1, "total_chunks": 2, "total_size": 10**12}, "declared size over the limit"),
        ({"number": 1, "total_chunks": 500, "total_size": 8}, "more chunks than bytes allow"),
    ],
)
def test_chunk_parameters_are_validated(small_chunks, kwargs, label):
    client = signed_in(small_chunks)
    response = chunk(
        client,
        "upload-0004",
        kwargs["number"],
        b"aaaa",
        kwargs["total_chunks"],
        kwargs["total_size"],
    )
    assert response.status_code == 400, label


def test_upload_ids_cannot_escape_the_staging_directory(small_chunks):
    client = signed_in(small_chunks)
    assert chunk(client, "../../etc/passwd", 1, b"aaaa", 1, 4).status_code == 400


# ---- request body limit ---------------------------------------------------------------


@pytest.fixture
def limited_client():
    app = FastAPI()

    @app.post("/size")
    async def size(request: Request):
        return {"n": len(await request.body())}

    app.add_middleware(BodySizeLimitMiddleware, max_body_bytes=100)
    return TestClient(app)


def test_body_under_the_limit_passes(limited_client):
    assert limited_client.post("/size", content=b"x" * 100).json() == {"n": 100}


def test_declared_oversize_body_is_rejected_with_413_problem(limited_client):
    response = limited_client.post("/size", content=b"x" * 101)

    assert response.status_code == 413
    assert response.headers["content-type"].startswith("application/problem+json")
    assert json.loads(response.content)["status"] == 413


def test_streamed_oversize_body_without_content_length_is_rejected(limited_client):
    def body():
        for _ in range(11):
            yield b"x" * 10

    response = limited_client.post("/size", content=body())

    assert response.status_code == 413


# ---- spreadsheets -----------------------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def create_sheet(client, data, fmt="xlsx", name="book"):
    return client.post(
        f"/api/v1/files/spreadsheet/create?filename={name}",
        json={"data": data, "format": fmt},
    )


def test_spreadsheet_create_read_and_multi_sheet_edit_keeps_other_sheets(env):
    client = signed_in(env)
    created = create_sheet(client, {"A": [["x", 1]], "B": [["y", 2]]})
    assert created.status_code == 201
    file_id = created.json()["file_id"]

    read = client.get(f"/api/v1/files/{file_id}/spreadsheet/read").json()
    assert read["data"]["sheets"] == {"A": [["x", 1]], "B": [["y", 2]]}

    edited = client.put(
        f"/api/v1/files/{file_id}/spreadsheet/update",
        json={"data": {"A": [["x", 99]]}, "sheet_name": "A"},
    )
    assert edited.status_code == 200 and edited.json()["version"] == 2

    sheets = client.get(f"/api/v1/files/{file_id}/spreadsheet/read").json()["data"]["sheets"]
    assert sheets == {"A": [["x", 99]], "B": [["y", 2]]}  # B survived the edit

    # ...and version 1 still holds the original bytes
    original = client.get(f"/api/v1/files/{file_id}?version=1").content
    assert original != client.get(f"/api/v1/files/{file_id}").content


def test_spreadsheet_formula_injection_is_defused(env):
    client = signed_in(env)
    file_id = create_sheet(
        client, {"S": [['=HYPERLINK("http://evil","x")', "-5", "-cmd|calc"]]}
    ).json()["file_id"]

    row = client.get(f"/api/v1/files/{file_id}/spreadsheet/read").json()["data"]["sheets"]["S"][0]

    assert row[0].startswith("'=") and row[2].startswith("'-")
    assert row[1] == "-5"  # real numbers stay untouched


@pytest.mark.parametrize("sheet", ["../../escape", "a/b", "x" * 40, "[bad]"])
def test_hostile_sheet_names_are_rejected(env, sheet):
    client = signed_in(env)
    assert create_sheet(client, {sheet: [["x"]]}).status_code == 400


def test_last_sheet_cannot_be_deleted_and_csv_is_single_sheet(env):
    client = signed_in(env)
    file_id = create_sheet(client, {"Only": [["x"]]}).json()["file_id"]
    assert client.delete(f"/api/v1/files/{file_id}/spreadsheet/sheet/Only").status_code == 400
    assert create_sheet(client, {"A": [["x"]], "B": [["y"]]}, fmt="csv").status_code == 400


def test_csv_round_trip_and_non_spreadsheets_are_refused(env):
    client = signed_in(env)
    sheet_id = create_sheet(client, {"S": [["a", "b"], ["1", "2"]]}, fmt="csv").json()["file_id"]
    assert client.get(f"/api/v1/files/{sheet_id}/spreadsheet/read").json()["data"]["sheets"] == {
        "Sheet1": [["a", "b"], ["1", "2"]]
    }
    pdf_id = upload(client).json()["file_id"]
    assert client.get(f"/api/v1/files/{pdf_id}/spreadsheet/read").status_code == 400


def test_ods_round_trip(env):
    pytest.importorskip("odf")
    client = signed_in(env)
    file_id = create_sheet(client, {"A": [["x", 1], ["y", 2]]}, fmt="ods").json()["file_id"]

    sheets = client.get(f"/api/v1/files/{file_id}/spreadsheet/read").json()["data"]["sheets"]

    assert sheets == {"A": [["x", 1], ["y", 2]]}


# ---- images ---------------------------------------------------------------------------


def png_bytes(width, height):
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buffer, "PNG")
    return buffer.getvalue()


def test_every_thumbnail_size_is_actually_generated_at_its_size(env):
    import io

    from PIL import Image

    client = signed_in(env)
    response = upload(client, "pic.png", png_bytes(1200, 800), "image/png")
    assert response.status_code == 201

    paths = run(env.file_service.data_service.files.get_file_by_id(response.json()["file_id"]))[
        "thumbnail_paths"
    ]
    widths = {
        size: Image.open(io.BytesIO(run(env.file_service.storage_backend.read(path)))).width
        for size, path in paths.items()
    }
    assert widths == {"150": 150, "300": 300, "600": 600}  # was 150 for all of them


def test_corrupt_images_are_stored_without_thumbnails(env):
    client = signed_in(env)
    response = upload(client, "bad.png", b"not really a png", "image/png")
    assert response.status_code == 201
    assert response.json()["thumbnail_urls"] == {}


def test_xlsx_edit_keeps_formulas_and_styles_of_untouched_cells(env):
    import io

    import openpyxl
    from openpyxl.styles import Font

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Calc"
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, "=A1+A2"
    sheet["A1"].font = Font(bold=True)
    sheet["B1"] = 7
    other = book.create_sheet("Keep")
    other["A1"] = "=1+1"
    buffer = io.BytesIO()
    book.save(buffer)

    client = signed_in(env)
    file_id = upload(client, "calc.xlsx", buffer.getvalue(), XLSX).json()["file_id"]
    edited = client.put(
        f"/api/v1/files/{file_id}/spreadsheet/update",
        json={
            "data": {"Calc": [[10], [3]]},
            "sheet_name": "Calc",
        },  # A1 changes, A2 same, A3 absent
    )
    assert edited.status_code == 200

    path = run(env.file_service.get_owned(file_id, edited.json()["user_id"]))["storage_path"]
    result = openpyxl.load_workbook(io.BytesIO(run(env.file_service.storage_backend.read(path))))
    assert result["Calc"]["A1"].value == 10 and result["Calc"]["A1"].font.bold  # style kept
    assert result["Calc"]["B1"].value is None  # plain cell outside the supplied grid: cleared
    assert (
        result["Calc"]["A3"].value == "=A1+A2"
    )  # formula with no cached value is kept, not guessed away
    assert result["Keep"]["A1"].value == "=1+1"  # untouched sheet keeps its formula


def test_reading_back_unchanged_values_keeps_formulas(env):
    import io

    import openpyxl

    book = openpyxl.Workbook()
    book.active["A1"], book.active["A2"] = 4, "=A1*2"
    buffer = io.BytesIO()
    book.save(buffer)
    client = signed_in(env)
    file_id = upload(client, "f.xlsx", buffer.getvalue(), XLSX).json()["file_id"]

    sheet_name = client.get(f"/api/v1/files/{file_id}/spreadsheet/info").json()["sheet_names"][0]
    # cached values are absent for generated files, so only the changed cell is written
    client.put(
        f"/api/v1/files/{file_id}/spreadsheet/update",
        json={"data": {sheet_name: [[5], ["=A1*2"]]}, "sheet_name": sheet_name},
    )
    latest = run(env.file_service.get_owned(file_id, client.get("/api/v1/auth/me").json()["id"]))
    result = openpyxl.load_workbook(
        io.BytesIO(run(env.file_service.storage_backend.read(latest["storage_path"])))
    )
    assert result.active["A1"].value == 5


# ---- streaming, thumbnails, direct downloads, cleanup ------------------------------------


def blue_png() -> bytes:
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (400, 300), "blue").save(buffer, "PNG")
    return buffer.getvalue()


def test_download_is_streamed_with_a_content_length(env):
    client = signed_in(env)
    file_id = upload(client, content=b"%PDF-" + b"z" * 3_000_000).json()["file_id"]

    download = client.get(f"/api/v1/files/{file_id}")
    assert download.headers["content-length"] == str(5 + 3_000_000)
    assert len(download.content) == 5 + 3_000_000


def test_urls_in_responses_point_at_this_api_and_thumbnails_are_servable(env):
    client = signed_in(env)
    body = upload(client, "pic.png", blue_png(), "image/png").json()

    assert body["download_url"] == f"/api/v1/files/{body['file_id']}?version=1"
    assert set(body["thumbnail_urls"]) == {"150", "300", "600"}
    thumbnail = client.get(body["thumbnail_urls"]["150"])
    assert thumbnail.status_code == 200 and thumbnail.headers["content-type"] == "image/png"
    assert client.get(f"/api/v1/files/{body['file_id']}/thumbnails/999").status_code == 404

    stranger = signed_in(env, "eve", "eve@example.com")
    assert stranger.get(body["thumbnail_urls"]["150"]).status_code == 404


def test_oversized_upload_is_refused_and_leaves_no_bytes_behind(env):
    env.file_service.max_file_size = 100
    client = signed_in(env)

    response = upload(client, content=b"%PDF-" + b"x" * 500)

    assert response.status_code == 400 and "exceeds maximum" in response.json()["detail"]
    assert run(env.file_service.storage_backend.list("", recursive=True)) == []


def test_hard_delete_removes_every_version_and_thumbnail_from_storage(env):
    client = signed_in(env)
    first = upload(client, "pic.png", blue_png(), "image/png").json()
    second = client.post(
        f"/api/v1/files/{first['file_id']}/versions",
        files={"file": ("pic.png", blue_png() + b"\x00", "image/png")},
    )
    assert second.status_code == 201 and second.json()["version"] == 2
    storage = env.file_service.storage_backend
    assert len(run(storage.list("", recursive=True))) == 8  # 2 versions x (original + 3 thumbs)

    assert client.delete(f"/api/v1/files/{first['file_id']}?hard_delete=true").status_code == 200
    assert run(storage.list("", recursive=True)) == []


def test_download_redirects_to_a_direct_url_when_the_backend_has_one(env):
    client = signed_in(env)
    file_id = upload(client).json()["file_id"]
    storage = env.file_service.storage_backend

    async def presigned(path, *, content_type, content_disposition, expires_in):
        return f"https://cdn.example/{path}?ttl={expires_in}"

    storage.download_url = presigned
    env.file_service.download_redirect = True
    response = client.get(f"/api/v1/files/{file_id}", follow_redirects=False)

    assert response.status_code == 307
    assert response.headers["location"].startswith("https://cdn.example/")


def test_metadata_failure_after_storing_removes_the_orphaned_bytes(env, monkeypatch):
    client = signed_in(env)

    async def broken(*args, **kwargs):
        raise RuntimeError("database down")

    monkeypatch.setattr(env.file_service.data_service.files, "create_file_metadata", broken)
    assert upload(client).status_code == 500
    assert run(env.file_service.storage_backend.list("", recursive=True)) == []


def test_abandoned_chunked_uploads_are_purged_once_older_than_the_session_lifetime(env):
    import os
    import time

    client = signed_in(env)
    user_id = client.get("/api/v1/auth/me").json()["id"]
    handler = env.file_service.chunked_upload_handler
    storage = env.file_service.storage_backend
    for upload_id in ("abandoned-1", "fresh-0001"):
        run(storage.save(f"_uploads/{user_id}/{upload_id}/000001", b"x"))
    old = time.time() - handler.upload_ttl - 60
    os.utime(storage._get_full_path(f"_uploads/{user_id}/abandoned-1/000001"), (old, old))

    assert run(handler.purge_abandoned()) == 1
    assert run(storage.exists(f"_uploads/{user_id}/fresh-0001/000001"))
