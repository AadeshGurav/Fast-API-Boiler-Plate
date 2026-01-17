"""Cross-platform Python-based spreadsheet processor using openpyxl and pandas."""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .base import BaseSpreadsheetProcessor

if TYPE_CHECKING:
    from app.core.interfaces.storage_interface import StorageInterface
    from app.services.logger import Logger
    from config import Config


class PythonSpreadsheetProcessor(BaseSpreadsheetProcessor):
    """Cross-platform Python-based spreadsheet processor for macOS and other platforms."""

    def __init__(
        self: PythonSpreadsheetProcessor,
        logger: Logger,
        config: Config,
        storage_backend: StorageInterface,
        *args: dict[str, Any],
        **kwargs: dict[str, Any],
    ) -> None:
        """Initialize Python spreadsheet processor.

        Args:
        ----
            logger: Logger instance.
            config: Configuration instance.
            storage_backend: Storage backend instance.
            *args: Additional arguments.
            **kwargs: Additional keyword arguments.

        """
        super().__init__(logger, config, storage_backend, *args, **kwargs)

        # Check for required libraries
        self._check_dependencies()

        self.logger.info(
            "PythonSpreadsheetProcessor initialized",
            extra={"service": "PythonSpreadsheetProcessor"},
        )

    def _check_dependencies(self: PythonSpreadsheetProcessor) -> None:
        """Check if required dependencies are available.

        Raises:
        ------
            ImportError: If required libraries are not installed.

        """
        try:
            import openpyxl  # noqa: F401
            import pandas  # noqa: F401
        except ImportError as e:
            raise ImportError(
                "openpyxl and pandas are required for spreadsheet processing. "
                "Install with: pip install openpyxl pandas"
            ) from e

    async def read_spreadsheet(
        self: PythonSpreadsheetProcessor,
        file_path: str,
        sheet_name: str | None = None,
    ) -> dict[str, list[list[Any]]]:
        """Read and parse spreadsheet data.

        Args:
        ----
            file_path: Path to spreadsheet file.
            sheet_name: Optional specific sheet to read.

        Returns:
        -------
            Dictionary mapping sheet names to data (list of lists).

        Raises:
        ------
            RuntimeError: If reading fails.

        """
        import openpyxl
        import pandas

        content = await self.storage_backend.read(file_path)
        if not content:
            raise RuntimeError(f"File not found: {file_path}")

        file_ext = Path(file_path).suffix.lower()

        with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp_file:
            tmp_path = Path(tmp_file.name)
            tmp_path.write_bytes(content)

            try:
                result: dict[str, list[list[Any]]] = {}

                if file_ext == ".csv":
                    df = pandas.read_csv(tmp_path)
                    sheet_name_actual = sheet_name or "Sheet1"
                    result[sheet_name_actual] = df.values.tolist()
                    if not df.empty and df.columns.tolist():
                        result[sheet_name_actual] = [
                            df.columns.tolist()
                        ] + result[sheet_name_actual]

                elif file_ext in [".xlsx", ".xls"]:
                    if file_ext == ".xlsx":
                        workbook = openpyxl.load_workbook(tmp_path, data_only=True)
                    else:
                        df_dict = pandas.read_excel(tmp_path, sheet_name=None)
                        workbook = None
                        for name, df in df_dict.items():
                            if sheet_name and name != sheet_name:
                                continue
                            result[name] = df.values.tolist()
                            if not df.empty and df.columns.tolist():
                                result[name] = [df.columns.tolist()] + result[name]
                        return result

                    sheet_names = workbook.sheetnames

                    if sheet_name and sheet_name not in sheet_names:
                        raise RuntimeError(f"Sheet '{sheet_name}' not found in workbook")

                    for name in sheet_names:
                        if sheet_name and name != sheet_name:
                            continue

                        sheet = workbook[name]
                        rows = []
                        for row in sheet.iter_rows(values_only=True):
                            rows.append(list(row))

                        result[name] = rows

                elif file_ext == ".ods":
                    try:
                        df_dict = pandas.read_excel(tmp_path, sheet_name=None, engine="odf")
                    except ImportError:
                        raise RuntimeError(
                            "ODS file support requires odfpy. Install with: pip install odfpy"
                        )
                    for name, df in df_dict.items():
                        if sheet_name and name != sheet_name:
                            continue
                        result[name] = df.values.tolist()
                        if not df.empty and df.columns.tolist():
                            result[name] = [df.columns.tolist()] + result[name]

                else:
                    raise RuntimeError(f"Unsupported file format: {file_ext}")

                if sheet_name and sheet_name not in result:
                    raise RuntimeError(f"Sheet '{sheet_name}' not found in workbook")

                return result

            finally:
                tmp_path.unlink(missing_ok=True)

    async def create_spreadsheet(
        self: PythonSpreadsheetProcessor,
        file_path: str,
        data: dict[str, list[list[Any]]],
        format: str = "xlsx",
        headers: dict[str, list[str]] | None = None,
    ) -> bool:
        """Create new spreadsheet from data.

        Args:
        ----
            file_path: Output file path.
            data: Dictionary mapping sheet names to data (list of lists).
            format: Output format (xlsx, ods, csv).
            headers: Optional headers per sheet.

        Returns:
        -------
            True if successful, False otherwise.

        """
        import openpyxl
        import pandas

        if format == "csv" and len(data) > 1:
            self.logger.warning(
                "CSV format only supports single sheet, using first sheet",
                extra={"service": "PythonSpreadsheetProcessor"},
            )
            data = {list(data.keys())[0]: list(data.values())[0]}

        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{format}") as tmp_file:
            tmp_path = Path(tmp_file.name)

            try:
                if format == "csv":
                    sheet_name = list(data.keys())[0]
                    rows = data[sheet_name]
                    df = pandas.DataFrame(rows[1:] if rows else [], columns=rows[0] if rows else [])
                    if headers and sheet_name in headers:
                        df.columns = headers[sheet_name]
                    df.to_csv(tmp_path, index=False)

                elif format == "xlsx":
                    workbook = openpyxl.Workbook()
                    workbook.remove(workbook.active)

                    for sheet_name, rows in data.items():
                        sheet = workbook.create_sheet(title=sheet_name)

                        if headers and sheet_name in headers:
                            sheet.append(headers[sheet_name])

                        for row in rows:
                            sheet.append(row)

                    workbook.save(tmp_path)

                elif format == "ods":
                    try:
                        with pandas.ExcelWriter(tmp_path, engine="odf") as writer:
                            for sheet_name, rows in data.items():
                                df = pandas.DataFrame(
                                    rows[1:] if rows else [], columns=rows[0] if rows else []
                                )
                                if headers and sheet_name in headers:
                                    df.columns = headers[sheet_name]
                                df.to_excel(writer, sheet_name=sheet_name, index=False)
                    except ImportError:
                        raise RuntimeError(
                            "ODS file support requires odfpy. Install with: pip install odfpy"
                        )

                else:
                    self.logger.error(
                        f"Unsupported format: {format}",
                        extra={"service": "PythonSpreadsheetProcessor"},
                    )
                    return False

                output_content = tmp_path.read_bytes()
                saved = await self.storage_backend.save(
                    file_path, output_content, overwrite=True
                )

                if saved:
                    self.logger.info(
                        f"Created spreadsheet: {file_path}",
                        extra={"service": "PythonSpreadsheetProcessor"},
                    )

                return saved

            finally:
                tmp_path.unlink(missing_ok=True)

    async def update_spreadsheet(
        self: PythonSpreadsheetProcessor,
        file_path: str,
        data: dict[str, list[list[Any]]],
        sheet_name: str | None = None,
    ) -> bool:
        """Update existing spreadsheet.

        Args:
        ----
            file_path: Path to spreadsheet file.
            data: Dictionary mapping sheet names to data (list of lists).
            sheet_name: Optional specific sheet to update.

        Returns:
        -------
            True if successful, False otherwise.

        """
        existing_data = await self.read_spreadsheet(file_path)

        if sheet_name:
            if sheet_name not in existing_data:
                raise RuntimeError(f"Sheet '{sheet_name}' not found")
            existing_data[sheet_name] = data.get(sheet_name, existing_data[sheet_name])
        else:
            existing_data.update(data)

        content = await self.storage_backend.read(file_path)
        if not content:
            return False

        file_ext = Path(file_path).suffix.lower()
        format_map = {".xlsx": "xlsx", ".ods": "ods", ".csv": "csv", ".xls": "xlsx"}
        format_type = format_map.get(file_ext, "xlsx")

        return await self.create_spreadsheet(file_path, existing_data, format=format_type)

    async def delete_sheet(
        self: PythonSpreadsheetProcessor,
        file_path: str,
        sheet_name: str,
    ) -> bool:
        """Delete a sheet from workbook.

        Args:
        ----
            file_path: Path to spreadsheet file.
            sheet_name: Name of sheet to delete.

        Returns:
        -------
            True if successful, False otherwise.

        """
        existing_data = await self.read_spreadsheet(file_path)

        if sheet_name not in existing_data:
            raise RuntimeError(f"Sheet '{sheet_name}' not found")

        del existing_data[sheet_name]

        if not existing_data:
            raise RuntimeError("Cannot delete last sheet from workbook")

        content = await self.storage_backend.read(file_path)
        if not content:
            return False

        file_ext = Path(file_path).suffix.lower()
        format_map = {".xlsx": "xlsx", ".ods": "ods", ".csv": "csv", ".xls": "xlsx"}
        format_type = format_map.get(file_ext, "xlsx")

        return await self.create_spreadsheet(file_path, existing_data, format=format_type)

    async def get_sheet_info(
        self: PythonSpreadsheetProcessor,
        file_path: str,
    ) -> dict[str, Any]:
        """Get workbook/sheet metadata.

        Args:
        ----
            file_path: Path to spreadsheet file.

        Returns:
        -------
            Dictionary with workbook info (sheet names, row/column counts, etc.).

        """
        data = await self.read_spreadsheet(file_path)

        sheet_info: dict[str, dict[str, int]] = {}
        for sheet_name, rows in data.items():
            row_count = len(rows)
            col_count = max((len(row) for row in rows), default=0)
            sheet_info[sheet_name] = {
                "row_count": row_count,
                "column_count": col_count,
            }

        file_ext = Path(file_path).suffix.lower()
        format_map = {".xlsx": "xlsx", ".ods": "ods", ".csv": "csv", ".xls": "xlsx"}
        format_type = format_map.get(file_ext, "xlsx")

        return {
            "sheet_names": list(data.keys()),
            "sheet_info": sheet_info,
            "format": format_type,
            "total_sheets": len(data),
        }


__all__ = ["PythonSpreadsheetProcessor"]
