"""Spreadsheet reading and writing (xlsx, ods, csv) in pure Python.

One implementation on every platform: ``openpyxl`` for xlsx, ``odfpy`` (through pandas)
for ods, the ``csv`` module for csv. No office suite is launched as a subprocess.

Reads return computed values. xlsx edits patch the package in place (``xlsx_patch``):
formulas, styles, charts and images of everything not overwritten survive. ods/csv edits
rebuild the file from values. Callers edit a *copy* (see
``FileService.modify_as_new_version``), so the previous version is always intact.
"""

from __future__ import annotations

import asyncio
import csv
import io
import re
from pathlib import Path
from typing import TYPE_CHECKING, Any

from app.services.base_service import BaseService
from app.services.file import xlsx_package, xlsx_patch

if TYPE_CHECKING:
    from app.core.interfaces.storage_interface import StorageInterface
    from app.services.logger import Logger
    from config import Config

SPREADSHEET_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.oasis.opendocument.spreadsheet": "ods",
    "text/csv": "csv",
}
SUPPORTED_FORMATS = {"xlsx", "ods", "csv"}
CSV_SHEET_NAME = "Sheet1"
FORBIDDEN_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_NAME_LENGTH = 31  # Excel's limit; keeps names portable across formats
NUMBER_LIKE = re.compile(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?")

Rows = list[list[Any]]


def neutralize(value: Any) -> Any:
    """Defuse spreadsheet formula injection in text cells.

    A text cell starting with ``=``, ``@`` (or ``+``/``-`` followed by a non-number)
    would be executed as a formula when someone opens the file in Excel, so it is
    stored with a leading apostrophe. Real numbers and ordinary text pass untouched.
    """
    if isinstance(value, str) and value:
        first = value[0]
        if first in "=@" or (first in "+-" and not NUMBER_LIKE.fullmatch(value)):
            return "'" + value
    return value


def validate_sheet_name(name: str) -> str:
    """Return ``name`` if it is a legal sheet name in every supported format."""
    if not name.strip() or len(name) > MAX_SHEET_NAME_LENGTH or FORBIDDEN_SHEET_CHARS.search(name):
        raise ValueError(
            f"Invalid sheet name {name!r}: 1-{MAX_SHEET_NAME_LENGTH} characters "
            "without any of [ ] : * ? / \\"
        )
    return name


class SpreadsheetProcessor(BaseService):
    """Read, create and edit spreadsheet files held in a storage backend."""

    def __init__(
        self: SpreadsheetProcessor,
        logger: Logger,
        config: Config,
        storage_backend: StorageInterface,
        *args: dict[str, Any],
        **kwargs: dict[str, Any],
    ) -> None:
        """Initialize SpreadsheetProcessor.

        Args:
        ----
            logger: Logger instance.
            config: Configuration instance.
            storage_backend: Where spreadsheet files live.
            *args: Additional arguments.
            **kwargs: Additional keyword arguments.

        """
        super().__init__(config, logger, *args, **kwargs)
        self.storage_backend = storage_backend
        self.max_cells: int = config["file_spreadsheet_max_cells"]
        self.max_unpacked_bytes: int = config["file_max_size"] * 4  # zip-bomb guard
        self.logger.info(
            "SpreadsheetProcessor initialized", extra={"service": "SpreadsheetProcessor"}
        )

    # ---- type helpers ---------------------------------------------------------
    @staticmethod
    def _is_spreadsheet_type(content_type: str) -> bool:
        """Whether the MIME type is a spreadsheet format this service can process."""
        return content_type.split(";")[0].strip().lower() in SPREADSHEET_TYPES

    @staticmethod
    def _format_of(storage_path: str) -> str:
        fmt = Path(storage_path).suffix.lower().lstrip(".")
        if fmt not in SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported spreadsheet format '.{fmt}'; use xlsx, ods or csv")
        return fmt

    def _check_size(self, data: dict[str, Rows]) -> None:
        cells = sum(len(row) for rows in data.values() for row in rows)
        if cells > self.max_cells:
            raise ValueError(f"Spreadsheet exceeds the limit of {self.max_cells} cells")

    # ---- (de)serialisation: blocking, so always run in a worker thread -------------
    @staticmethod
    def _parse(content: bytes, fmt: str, max_cells: int) -> dict[str, Rows]:
        """Read every sheet, refusing workbooks that expand past ``max_cells`` cells.

        The cap counts cells as the reader yields them (openpyxl pads each row to the sheet's
        declared width), so a few-KB file claiming ``A1:XFD1048576`` cannot allocate millions
        of cells before the limit is checked.
        """
        cells = 0

        def take(rows: Any) -> Rows:
            nonlocal cells
            kept: Rows = []
            for row in rows:
                cells += len(row)
                if cells > max_cells:
                    raise ValueError(f"Spreadsheet exceeds the limit of {max_cells} cells")
                kept.append(list(row))
            return kept

        if fmt == "csv":
            text = content.decode("utf-8-sig")
            return {CSV_SHEET_NAME: take(csv.reader(io.StringIO(text)))}

        if fmt == "xlsx":
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                return {
                    sheet.title: take(sheet.iter_rows(values_only=True))
                    for sheet in workbook.worksheets
                }
            finally:
                workbook.close()

        import pandas

        frames = pandas.read_excel(io.BytesIO(content), sheet_name=None, engine="odf", header=None)
        return {
            name: take(frame.astype(object).where(frame.notna(), None).values.tolist())
            for name, frame in frames.items()
        }

    @staticmethod
    def _serialize(data: dict[str, Rows], fmt: str, headers: dict[str, list[str]] | None) -> bytes:
        def with_headers(name: str, rows: Rows) -> Rows:
            head = [headers[name]] if headers and name in headers else []
            return [[neutralize(cell) for cell in row] for row in head + rows]

        if fmt == "csv":
            name, rows = next(iter(data.items()))
            buffer = io.StringIO(newline="")
            csv.writer(buffer).writerows(with_headers(name, rows))
            return buffer.getvalue().encode("utf-8")

        if fmt == "xlsx":
            import openpyxl

            workbook = openpyxl.Workbook(write_only=True)
            for name, rows in data.items():
                sheet = workbook.create_sheet(title=name)
                for row in with_headers(name, rows):
                    sheet.append(row)
            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue()

        import pandas

        output = io.BytesIO()
        with pandas.ExcelWriter(output, engine="odf") as writer:
            for name, rows in data.items():
                pandas.DataFrame(with_headers(name, rows)).to_excel(
                    writer, sheet_name=name, index=False, header=False
                )
        return output.getvalue()

    @staticmethod
    def _diff(rows: Rows, current: Rows) -> xlsx_patch.Edits:
        """Cells where ``rows`` differs from the sheet's current computed values.

        Sending back what a read returned changes nothing, so formulas survive; cells
        beyond ``rows`` are cleared, matching "replace the rows of this sheet".
        """
        height = max(len(rows), len(current))
        width = max(
            max((len(r) for r in rows), default=0), max((len(r) for r in current), default=0)
        )
        edits: xlsx_patch.Edits = {}
        for r in range(height):
            for c in range(width):
                new = rows[r][c] if r < len(rows) and c < len(rows[r]) else None
                old = current[r][c] if r < len(current) and c < len(current[r]) else None
                if new != old:
                    edits[(r + 1, c + 1)] = neutralize(new)
        return edits

    async def _edit_in_place(
        self,
        storage_path: str,
        data: dict[str, Rows],
        current: dict[str, Rows],
        delete: str | None = None,
    ) -> bool:
        """Patch the xlsx package (see ``xlsx_patch``): charts, images, formulas and
        styles of everything not overwritten are preserved."""
        content = await self.storage_backend.read(storage_path)
        if content is None:
            raise ValueError("File not found in storage")
        for name in data:
            validate_sheet_name(name)
        self._check_size(data)
        edited = await asyncio.to_thread(
            xlsx_package.edit_xlsx,
            content,
            edits={n: self._diff(rows, current[n]) for n, rows in data.items() if n in current},
            new_sheets={n: rows for n, rows in data.items() if n not in current},
            delete=delete,
            max_unpacked_bytes=self.max_unpacked_bytes,
        )
        return await self.storage_backend.save(storage_path, edited, overwrite=True)

    async def _load(self, storage_path: str) -> dict[str, Rows]:
        fmt = self._format_of(storage_path)
        content = await self.storage_backend.read(storage_path)
        if content is None:
            raise ValueError("File not found in storage")
        try:
            return await asyncio.to_thread(self._parse, content, fmt, self.max_cells)
        except ImportError as e:
            raise ValueError(f"Reading .{fmt} files needs an optional library: {e}") from e

    async def _store(
        self, storage_path: str, data: dict[str, Rows], headers: dict[str, list[str]] | None = None
    ) -> bool:
        fmt = self._format_of(storage_path)
        if not data:
            raise ValueError("A spreadsheet needs at least one sheet")
        for name in data:
            validate_sheet_name(name)
        self._check_size(data)
        try:
            content = await asyncio.to_thread(self._serialize, data, fmt, headers)
        except ImportError as e:
            raise ValueError(f"Writing .{fmt} files needs an optional library: {e}") from e
        return await self.storage_backend.save(storage_path, content, overwrite=True)

    # ---- public API ---------------------------------------------------------------
    async def read_spreadsheet(
        self: SpreadsheetProcessor, storage_path: str, sheet_name: str | None = None
    ) -> dict[str, Rows]:
        """Cell values per sheet (all sheets, or just ``sheet_name``)."""
        data = await self._load(storage_path)
        if sheet_name is None:
            return data
        if sheet_name not in data:
            raise ValueError(f"Sheet '{sheet_name}' not found")
        return {sheet_name: data[sheet_name]}

    async def create_spreadsheet(
        self: SpreadsheetProcessor,
        storage_path: str,
        data: dict[str, Rows],
        format: str = "xlsx",
        headers: dict[str, list[str]] | None = None,
    ) -> bool:
        """Write a new spreadsheet; the format is taken from ``storage_path``'s extension."""
        if self._format_of(storage_path) != format:
            raise ValueError(f"storage path extension does not match format '{format}'")
        if format == "csv" and len(data) > 1:
            raise ValueError("CSV holds a single sheet")
        return await self._store(storage_path, data, headers)

    async def update_spreadsheet(
        self: SpreadsheetProcessor,
        storage_path: str,
        data: dict[str, Rows],
        sheet_name: str | None = None,
    ) -> bool:
        """Replace the rows of the given sheets (all in ``data``, or only ``sheet_name``);
        sheets not mentioned are kept as they are."""
        existing = await self._load(storage_path)
        if self._format_of(storage_path) == "xlsx":
            if sheet_name is not None:
                if sheet_name not in existing:
                    raise ValueError(f"Sheet '{sheet_name}' not found")
                if sheet_name not in data:
                    raise ValueError(f"No data supplied for sheet '{sheet_name}'")
                data = {sheet_name: data[sheet_name]}
            return await self._edit_in_place(storage_path, data, existing)
        if sheet_name is not None:
            if sheet_name not in existing:
                raise ValueError(f"Sheet '{sheet_name}' not found")
            if sheet_name not in data:
                raise ValueError(f"No data supplied for sheet '{sheet_name}'")
            existing[sheet_name] = data[sheet_name]
        else:
            existing.update(data)
        return await self._store(storage_path, existing)

    async def delete_sheet(self: SpreadsheetProcessor, storage_path: str, sheet_name: str) -> bool:
        """Remove one sheet; a workbook cannot lose its last sheet."""
        existing = await self._load(storage_path)
        if sheet_name not in existing:
            raise ValueError(f"Sheet '{sheet_name}' not found")
        if len(existing) == 1:
            raise ValueError("Cannot delete the last sheet of a workbook")
        if self._format_of(storage_path) == "xlsx":
            return await self._edit_in_place(storage_path, {}, existing, delete=sheet_name)
        del existing[sheet_name]
        return await self._store(storage_path, existing)

    async def get_sheet_info(self: SpreadsheetProcessor, storage_path: str) -> dict[str, Any]:
        """Sheet names and row/column counts."""
        data = await self._load(storage_path)
        return {
            "sheet_names": list(data),
            "sheet_info": {
                name: {
                    "row_count": len(rows),
                    "column_count": max((len(row) for row in rows), default=0),
                }
                for name, rows in data.items()
            },
            "format": self._format_of(storage_path),
            "total_sheets": len(data),
        }


__all__ = ["SpreadsheetProcessor", "neutralize", "validate_sheet_name"]
